import io
import csv
import json

SUPPORTED_EXTENSIONS = {".txt", ".md", ".pdf", ".docx", ".csv", ".json", ".html", ".htm", ".xlsx"}


def extract_text(content: bytes, filename: str) -> str:
    ext = _ext(filename)
    if ext in (".txt", ".md"):
        return content.decode("utf-8")
    if ext == ".pdf":
        return _from_pdf(content)
    if ext == ".docx":
        return _from_docx(content)
    if ext == ".csv":
        return _from_csv(content)
    if ext == ".json":
        return _from_json(content)
    if ext in (".html", ".htm"):
        return _from_html(content)
    if ext == ".xlsx":
        return _from_xlsx(content)
    raise ValueError(f"Formato non supportato: {ext}")


def _ext(filename: str) -> str:
    dot = filename.rfind(".")
    return filename[dot:].lower() if dot != -1 else ""


def _from_pdf(content: bytes) -> str:
    from pypdf import PdfReader
    reader = PdfReader(io.BytesIO(content))
    pages = [page.extract_text() or "" for page in reader.pages]
    return "\n\n".join(p.strip() for p in pages if p.strip())


def _from_docx(content: bytes) -> str:
    from docx import Document
    doc = Document(io.BytesIO(content))
    paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
    return "\n\n".join(paragraphs)


def _from_csv(content: bytes) -> str:
    text = content.decode("utf-8", errors="replace")
    reader = csv.reader(io.StringIO(text))
    rows = [", ".join(row) for row in reader if any(cell.strip() for cell in row)]
    return "\n".join(rows)


def _from_json(content: bytes) -> str:
    data = json.loads(content.decode("utf-8"))
    return json.dumps(data, ensure_ascii=False, indent=2)


def _from_xlsx(content: bytes) -> str:
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True)
    parts = []
    for sheet in wb.worksheets:
        rows = []
        for row in sheet.iter_rows(values_only=True):
            cells = [str(c) if c is not None else "" for c in row]
            if any(c.strip() for c in cells):
                rows.append("\t".join(cells))
        if rows:
            parts.append(f"[Foglio: {sheet.title}]\n" + "\n".join(rows))
    return "\n\n".join(parts)


def _from_html(content: bytes) -> str:
    from bs4 import BeautifulSoup
    soup = BeautifulSoup(content, "lxml")
    for tag in soup(["script", "style", "nav", "footer", "header"]):
        tag.decompose()
    return soup.get_text(separator="\n", strip=True)
